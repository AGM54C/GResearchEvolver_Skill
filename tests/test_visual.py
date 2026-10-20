"""Artifact and human-gate tests; fixtures do not measure scientific taste."""
import copy
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET

REPO = Path(__file__).resolve().parents[1]
SKILLS = REPO / 'skills'
sys.path.insert(0, str(SKILLS / '_shared'))
sys.path.insert(0, str(SKILLS / '_shared/runtime'))
from visual_artifacts import VisualError, load_spec, sha256, write_json, read_json
from gresearch.core import Store, ResearchError
from gresearch.execution import run_workflow
from gresearch import hci


def module(name, relative):
    spec = importlib.util.spec_from_file_location(name, SKILLS / relative)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


example = module('visual_example', '_shared/examples/visual_research.py')
figures = module('visual_figures', 'figure-mimic/scripts/build_figure.py')
diagrams = module('visual_diagrams', 'figure-mimic/scripts/build_diagram.py')
decks = module('visual_decks', 'slides-builder/scripts/build_deck.py')
editing = module('visual_editing', 'slides-builder/scripts/deck_tools.py')


class VisualTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.sources = example.materials(self.root / 'sources')

    def tearDown(self):
        self.temp.cleanup()

    def spec(self, kind):
        return read_json(self.sources / (kind + '-spec.json'))

    def change(self, kind, spec):
        path = self.sources / (kind + '-spec.json')
        write_json(path, spec)
        return path

    def test_source_drift_refuses_build_before_publication(self):
        (self.sources / 'sample.csv').write_text('changed', encoding='utf-8')
        for builder, kind in ((figures, 'figure'), (diagrams, 'diagram'), (decks, 'deck')):
            with self.assertRaises(VisualError):
                builder.build(self.sources / (kind + '-spec.json'), self.root / kind)
            self.assertFalse((self.root / kind).exists())

    def test_figure_outputs_have_expected_dimensions_and_rebuildable_values(self):
        import fitz
        from PIL import Image
        out = self.root / 'figure'
        record = figures.build(self.sources / 'figure-spec.json', out)
        with fitz.open(out / 'figure.pdf') as pdf:
            self.assertAlmostEqual(pdf[0].rect.width, 180 / 25.4 * 72, places=2)
            self.assertAlmostEqual(pdf[0].rect.height, 100 / 25.4 * 72, places=2)
            self.assertIn('Supplied intervals', pdf[0].get_text())
        with Image.open(out / 'figure.png') as png:
            self.assertAlmostEqual(png.width, 180 / 25.4 * 160, delta=1)
        self.assertEqual(record['sources'][0]['sha256'], sha256(self.sources / 'sample.csv'))
        self.assertEqual(record['human_review'], 'pending')
        saved = read_json(out / 'spec.json')
        self.assertEqual(saved['panels'], self.spec('figure')['panels'])
        self.assertEqual(sha256(out / saved['sources'][0]['path']), sha256(self.sources / 'sample.csv'))
        # The output remains rebuildable after the original input disappears.
        (self.sources / 'sample.csv').unlink()
        figures.build(out / 'spec.json', self.root / 'rebuilt')
        before = sha256(out / 'figure.png')
        with self.assertRaises(VisualError): figures.build(out / 'spec.json', out)
        self.assertEqual(before, sha256(out / 'figure.png'))

    def test_missing_line_observation_is_a_gap_in_exported_vector_path(self):
        csv = self.sources / 'sample.csv'
        csv.write_text(csv.read_text().replace('2,3,2.4,3.6,2', '2,,2.4,3.6,2'), encoding='utf-8')
        spec = self.spec('figure')
        spec['sources'][0]['sha256'] = sha256(csv)
        out = self.root / 'missing'
        figures.build(self.change('figure', spec), out)
        root = ET.parse(out / 'figure.svg').getroot()
        paths = [p.get('d', '') for p in root.iter('{http://www.w3.org/2000/svg}path') if '#0072b2' in p.get('style', '')]
        self.assertTrue(any(p.count('M') >= 2 for p in paths), paths)

    def test_bad_interval_and_log_values_do_not_become_success(self):
        spec = self.spec('figure')
        spec['panels'][0]['series'][0]['interval']['low'] = 'upper'
        spec['panels'][0]['series'][0]['interval']['high'] = 'lower'
        with self.assertRaises(VisualError): figures.build(self.change('figure', spec), self.root / 'bad')
        spec['panels'][0]['series'][0].pop('interval')
        spec['panels'][0]['yscale'] = 'log'
        csv = self.sources / 'sample.csv'
        csv.write_text(csv.read_text().replace('1,2,', '1,-2,'), encoding='utf-8')
        spec['sources'][0]['sha256'] = sha256(csv)
        with self.assertRaises(VisualError): figures.build(self.change('figure', spec), self.root / 'bad-log')

    def test_bounds_screen_rejects_actual_clipped_title(self):
        spec = self.spec('figure')
        spec['panels'][0]['title'] = 'Very long title ' * 40
        with self.assertRaisesRegex(VisualError, 'Text outside figure bounds'):
            figures.build(self.change('figure', spec), self.root / 'clipped')
        self.assertFalse((self.root / 'clipped').exists())

    def test_distribution_and_heatmap_exports_preserve_observations(self):
        import fitz
        spec = self.spec('figure')
        spec['size_mm'] = [200, 160]
        common = {'source': 'data', 'xlabel': 'Fixture group', 'ylabel': 'Fixture value', 'claims': ['demo']}
        series = [{'y': 'value', 'label': 'Fixture'}]
        spec['panels'] = [{**common, 'kind': kind, 'title': kind, 'series': series,
                           **({'bins': [0, 1, 2, 3, 4, 5]} if kind == 'histogram' else {})}
                          for kind in ('histogram', 'box', 'violin')]
        spec['panels'].append({**common, 'kind': 'heatmap', 'title': 'heatmap',
                               'columns': ['value', 'comparison'], 'colorbar': 'Supplied value'})
        out = self.root / 'distributions'
        record = figures.build(self.change('figure', spec), out)
        self.assertTrue(all(p['rows'] == 4 and p['missing_y_cells'] == 0 for p in record['checks']['panels']))
        with fitz.open(out / 'figure.pdf') as document:
            self.assertTrue(all(kind in document[0].get_text() for kind in ('histogram', 'box', 'violin', 'heatmap')))

    def test_diagram_preserves_editable_nodes_edges_and_rejects_dangling_links(self):
        import fitz
        out = self.root / 'diagram'
        diagrams.build(self.sources / 'diagram-spec.json', out)
        root = ET.parse(out / 'diagram.drawio').getroot()
        nodes = [c for c in root.iter('mxCell') if c.get('vertex') == '1']
        edges = [c for c in root.iter('mxCell') if c.get('edge') == '1']
        self.assertEqual({n.get('id') for n in nodes}, {'brief', 'build', 'review'})
        self.assertEqual({(e.get('source'), e.get('target')) for e in edges}, {('brief', 'build'), ('build', 'review')})
        # Directed relations must retain visible arrowheads after SVG-to-PDF.
        with fitz.open(out / 'diagram.pdf') as pdf:
            arrowheads = [p for p in pdf[0].get_drawings() if p['fill'] is not None
                          and abs(p['rect'].width - 12) < .1 and abs(p['rect'].height - 10) < .1]
            self.assertEqual(len(arrowheads), 2)
        spec = self.spec('diagram'); spec['edges'][0]['to'] = 'missing'
        with self.assertRaises(VisualError): diagrams.build(self.change('diagram', spec), self.root / 'bad-diagram')

    def test_pptx_contains_native_chart_exact_values_and_notes(self):
        from pptx import Presentation
        out = self.root / 'deck'
        record = decks.build(self.sources / 'deck-spec.json', out)
        deck = Presentation(out / 'slides.pptx')
        self.assertEqual(len(deck.slides), 4)
        chart = next(shape.chart for shape in deck.slides[1].shapes if shape.has_chart)
        self.assertEqual(tuple(chart.series[0].values), (2, 3, 2.5, 4))
        self.assertTrue(all(s.notes_slide.notes_text_frame.text.strip() for s in deck.slides))
        self.assertEqual(record['checks']['office_render'], 'not_run')
        self.assertEqual(record['human_review'], 'pending')
        self.assertFalse(any(shape['outside_canvas'] for s in editing.inspect(out / 'slides.pptx')['slides'] for shape in s['shapes']))

    def test_deck_timing_and_overflow_refuse_to_publish(self):
        spec = self.spec('deck'); spec['duration_seconds'] = 10
        with self.assertRaises(VisualError): decks.build(self.change('deck', spec), self.root / 'timing')
        spec['duration_seconds'] = 240; spec['slides'][0]['title'] = 'Overlong title ' * 80
        with self.assertRaises(VisualError): decks.build(self.change('deck', spec), self.root / 'overflow')

    def test_native_chart_does_not_silently_drop_uncertainty_or_missing_values(self):
        spec = self.spec('deck')
        spec['slides'][1]['chart']['series'][0]['interval'] = {'low': 'lower', 'high': 'upper'}
        with self.assertRaises(VisualError): decks.build(self.change('deck', spec), self.root / 'interval')
        spec['slides'][1]['chart']['series'][0].pop('interval')
        csv = self.sources / 'sample.csv'; csv.write_text(csv.read_text().replace('2,3,', '2,,'), encoding='utf-8')
        spec['sources'][0]['sha256'] = sha256(csv)
        with self.assertRaises(VisualError): decks.build(self.change('deck', spec), self.root / 'missing-chart')

    def test_template_edit_preserves_other_parts_and_rejects_stale_text(self):
        from pptx import Presentation
        out = self.root / 'deck'; decks.build(self.sources / 'deck-spec.json', out)
        source = out / 'slides.pptx'; original = sha256(source)
        inventory = editing.inspect(source)
        title = inventory['slides'][0]['shapes'][0]
        spec = {**self.spec('deck'), 'deck_source': 'template',
                'sources': [{'id': 'template', 'path': str(source), 'sha256': original}], 'claims': [],
                'edits': [{'part': inventory['slides'][0]['part'], 'shape_id': title['id'], 'before': title['text'], 'after': 'A reviewed framing change'}]}
        path = self.change('edit', spec)
        revised = self.root / 'revised'; editing.edit(path, revised)
        self.assertEqual(original, sha256(source))
        with zipfile.ZipFile(source) as a, zipfile.ZipFile(revised / 'slides.pptx') as b:
            self.assertEqual(a.namelist(), b.namelist())
            for name in a.namelist():
                if name != spec['edits'][0]['part']: self.assertEqual(a.read(name), b.read(name), name)
        self.assertEqual(Presentation(revised / 'slides.pptx').slides[0].shapes[0].text, 'A reviewed framing change')
        spec['edits'][0]['before'] = 'stale'
        with self.assertRaises(VisualError): editing.edit(self.change('edit', spec), self.root / 'stale')

    def test_exact_quote_is_checked_but_never_called_semantic_verification(self):
        source = self.sources / 'source.txt'; source.write_text('A reported observation.', encoding='utf-8')
        spec = self.spec('figure')
        spec['sources'].append({'id': 'text', 'path': 'source.txt', 'sha256': sha256(source)})
        spec['claims'] = [{'id': 'x', 'text': 'An interpretation', 'status': 'source_reported', 'evidence': [
            {'source': 'text', 'locator': 'first sentence', 'start': 0, 'end': 1, 'quote': 'B'}]}]
        with self.assertRaises(VisualError): load_spec(self.change('figure', spec))

    def test_real_visual_stages_wait_for_researcher_and_respect_revision(self):
        result = example.initialize(self.root / 'workflow')
        self.assertEqual(result['status'], 'awaiting_review')
        store = Store(self.root / 'workflow/state')
        try:
            run = result['run']
            self.assertFalse((store.work(run) / 'figure/figure.png').exists())
            for stage in ('brief', 'figure', 'diagram'):
                report = hci.read(store, run, stage)
                hci.feedback(store, run, stage, report['receipt']['id'], 'researcher', 'approve', 'SCRIPTED TEST ONLY; not human study data')
                state = run_workflow(store, run)
                self.assertEqual(state['status'], 'awaiting_review', state.get('blocker'))
            self.assertTrue((store.work(run) / 'deck/slides.pptx').exists())
            report = hci.read(store, run, 'deck')
            feedback = hci.feedback(store, run, 'deck', report['receipt']['id'], 'researcher', 'revise', 'SCRIPTED: revise the framing before adoption')
            with self.assertRaises(ResearchError): store.control(run, 'resume')
            self.assertNotEqual(store.get(run)[1]['status'], 'completed')
            self.assertEqual(feedback['feedback']['decision'], 'revise')
            self.assertTrue(store.verify()['valid'])
        finally:
            store.close()

if __name__ == '__main__':
    unittest.main()
