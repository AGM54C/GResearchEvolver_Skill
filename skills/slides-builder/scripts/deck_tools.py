"""Inspect or make exact text edits to an existing OOXML presentation copy."""
import argparse
from pathlib import Path
import sys
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / '_shared'))
from visual_artifacts import VisualError, finalize, load_spec, output_bundle, preview_pdf, source_path, write_json

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def inspect(path):
    from pptx import Presentation
    from pptx.oxml.ns import qn
    presentation = Presentation(path)
    pages = []
    for index, slide in enumerate(presentation.slides):
        shapes = []
        for shape in slide.shapes:
            outside = shape.left < 0 or shape.top < 0 or shape.left + shape.width > presentation.slide_width + 10 or shape.top + shape.height > presentation.slide_height + 10
            shapes.append({'id': shape.shape_id, 'name': shape.name, 'type': str(shape.shape_type),
                           'text': shape.text if shape.has_text_frame else None, 'outside_canvas': outside,
                           'text_runs': [node.text or '' for node in shape._element.iter(qn('a:t'))],
                           'native_chart': shape.has_chart})
        pages.append({'slide': index + 1, 'part': str(slide.part.partname).lstrip('/'), 'shapes': shapes,
                      'notes': slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''})
    return {'slides': pages, 'size_emu': [presentation.slide_width, presentation.slide_height],
            'scope': 'Object inventory and bounds only; not font rendering, overlap or scientific verification'}


def edit(spec_path, destination):
    from lxml import etree
    spec, sources, _ = load_spec(spec_path)
    source = source_path(sources, spec.get('deck_source'))
    edits = spec.get('edits', [])
    if not edits:
        raise VisualError('Provide exact edits from the presentation inventory')
    changes = {}
    with zipfile.ZipFile(source) as archive:
        if len(archive.namelist()) != len(set(archive.namelist())):
            raise VisualError('Duplicate ZIP entries')
        if sum(i.file_size for i in archive.infolist()) > 500000000:
            raise VisualError('Presentation exceeds the 500 MB uncompressed editing limit')
        for change in edits:
            part = change.get('part', '')
            if not part.startswith('ppt/slides/slide') or not part.endswith('.xml') or part not in archive.namelist():
                raise VisualError('Edit must target an existing slide XML part')
            root = changes.get(part)
            if root is None:
                root = etree.fromstring(archive.read(part), etree.XMLParser(resolve_entities=False, no_network=True))
            found = root.xpath('.//p:sp[p:nvSpPr/p:cNvPr/@id=$identity] | .//p:graphicFrame[p:nvGraphicFramePr/p:cNvPr/@id=$identity]',
                               namespaces=NS, identity=str(change['shape_id']))
            if len(found) != 1:
                raise VisualError('Exact shape id is missing or ambiguous')
            runs = found[0].xpath('.//a:t', namespaces=NS)
            run_index = change.get('run', 0)
            if type(run_index) is not int or not 0 <= run_index < len(runs) or runs[run_index].text != change.get('before'):
                raise VisualError('Stale or ambiguous text edit; inspect the original runs')
            if not isinstance(change.get('after'), str):
                raise VisualError('Replacement must be text')
            runs[run_index].text = change['after']
            if change['after'] != change['after'].strip():
                runs[run_index].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            changes[part] = root
        with output_bundle(destination) as out:
            filename = 'slides' + source.suffix.lower()
            with zipfile.ZipFile(out / filename, 'w') as result:
                for info in archive.infolist():
                    data = etree.tostring(changes[info.filename], encoding='UTF-8', xml_declaration=True) if info.filename in changes else archive.read(info.filename)
                    result.writestr(info, data)
            return finalize(out, spec, sources, {'edited_parts': sorted(changes), 'edits': len(edits),
                            'unchanged_parts': 'original uncompressed bytes retained', 'office_render': 'not_run',
                            'visual_review': 'pending; replacement text can change wrapping'}, builder=Path(__file__))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='operation', required=True)
    for command in ('inspect', 'edit', 'preview-pdf'):
        child = commands.add_parser(command)
        child.add_argument('input'); child.add_argument('--output', required=True)
    args = parser.parse_args()
    try:
        if args.operation == 'inspect':
            if Path(args.output).exists(): raise VisualError('Inventory output already exists')
            write_json(args.output, inspect(args.input))
        elif args.operation == 'edit':
            edit(args.input, args.output)
        else:
            with output_bundle(args.output) as out:
                write_json(out / 'render.json', preview_pdf(args.input, out / 'pages'))
    except (VisualError, OSError, KeyError, TypeError) as exc:
        parser.exit(2, str(exc) + '\n')
