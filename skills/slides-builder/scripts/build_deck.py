"""Build an editable, source-bound PPTX with notes and a slide evidence ledger."""
import argparse
import math
from pathlib import Path
import sys
import unicodedata

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / '_shared'))
from visual_artifacts import (VisualError, column, finalize, font_for, load_spec, nonempty,
                              number, output_bundle, references, source_path, table, write_json)


def text_box(slide, text, box, font, size=24, color='172638', bold=False):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    x, y, w, h = box
    # Conservative layout screen; actual Office rendering is still required.
    capacity = max(1, w * 72 / size)
    lines = sum(max(1, math.ceil(sum(1 if unicodedata.east_asian_width(c) in 'WF' else .56 for c in line) / capacity)) for line in text.split('\n'))
    if lines * size * 1.25 > h * 72:
        raise VisualError('Text may overflow; shorten or split the slide: ' + text[:65])
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(text.split('\n')):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name, paragraph.font.size = font, Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(6)
    return shape


def picture(slide, path, box, description):
    from PIL import Image
    from pptx.util import Inches
    with Image.open(path) as asset:
        w, h = asset.size
    x, y, bw, bh = box
    scale = min(bw / w, bh / h)
    shape = slide.shapes.add_picture(str(path), Inches(x + (bw - w * scale) / 2), Inches(y + (bh - h * scale) / 2),
                                     width=Inches(w * scale), height=Inches(h * scale))
    shape._element.nvPicPr.cNvPr.set('descr', nonempty(description, 'image alt text'))


def native_chart(slide, item, sources, box, font):
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    kind = item.get('kind')
    rows = table(sources, item.get('source'))
    if kind == 'scatter':
        data = XyChartData()
        for s in item['series']:
            series = data.add_series(nonempty(s.get('label'), 'series label'))
            for x, y in zip(column(rows, item['x']), column(rows, s['y'])):
                if not math.isfinite(x) or not math.isfinite(y):
                    raise VisualError('Native chart has missing values; use an explicitly rendered missing-data figure')
                series.add_data_point(x, y)
        chart_type = XL_CHART_TYPE.XY_SCATTER
    elif kind in {'bar', 'line'}:
        data = CategoryChartData()
        data.categories = column(rows, item['x'], numeric=False)
        for s in item['series']:
            values = column(rows, s['y'])
            if any(not math.isfinite(v) for v in values):
                raise VisualError('Native chart has missing values; use an explicitly rendered missing-data figure')
            data.add_series(nonempty(s.get('label'), 'series label'), values)
        chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED if kind == 'bar' else XL_CHART_TYPE.LINE_MARKERS
    else:
        raise VisualError('Native chart kind must be bar, line or scatter')
    if any(s.get('interval') for s in item['series']):
        raise VisualError('Use a source-rendered figure to preserve uncertainty intervals')
    chart = slide.shapes.add_chart(chart_type, *(Inches(v) for v in box), data).chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.name, chart.legend.font.size = font, Pt(14)
    chart.plots[0].vary_by_categories = False
    for index, series in enumerate(chart.series):
        color = RGBColor.from_string(('0072B2', 'D55E00', '009E73', 'CC79A7')[index % 4])
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
        series.format.line.color.rgb = color
    for axis, label in ((chart.category_axis, item.get('xlabel')), (chart.value_axis, item.get('ylabel'))):
        axis.has_title = True
        axis.axis_title.text_frame.text = nonempty(label, 'chart axis label')
        for paragraph in axis.axis_title.text_frame.paragraphs:
            paragraph.font.name, paragraph.font.size = font, Pt(15)
        axis.tick_labels.font.name, axis.tick_labels.font.size = font, Pt(15)
    if kind == 'bar':
        all_values = [v for s in item['series'] for v in column(rows, s['y'])]
        if min(all_values) >= 0: chart.value_axis.minimum_scale = 0
        elif max(all_values) <= 0: chart.value_axis.maximum_scale = 0


def build(spec_path, destination):
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.dml.color import RGBColor
    spec, sources, claims = load_spec(spec_path)
    slides = spec.get('slides', [])
    if not 1 <= len(slides) <= 200:
        raise VisualError('Provide 1..200 slides')
    all_text = str(slides) + str(spec['sources']) + str(spec.get('claims', []))
    font = font_for(all_text, spec.get('font'))
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = spec['purpose']
    deck.core_properties.subject = spec['evidence_status']
    deck.core_properties.author = spec.get('presenter', '')
    deck.core_properties.keywords = 'evidence-bound; researcher review pending'
    ledger, total_seconds = [], 0
    for index, item in enumerate(slides):
        title = nonempty(item.get('title'), 'slide title')
        notes = nonempty(item.get('notes'), 'speaker notes')
        claim_ids = references(item, claims)
        duration = number(item.get('seconds'), 'slide seconds', 0, 1800)
        if not item.get('backup', False): total_seconds += duration
        layout = item.get('layout', 'evidence')
        if layout not in {'title', 'evidence', 'comparison', 'figure', 'discussion'}:
            raise VisualError('Unsupported slide layout')
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
        text_box(slide, title, (.6, .35, 12.1, 1.0), font, size=30, bold=True)
        assets = []
        body = item.get('body', [])
        if not isinstance(body, list) or any(not isinstance(t, str) for t in body):
            raise VisualError('body must be a list of text paragraphs')
        if layout == 'comparison':
            columns = item.get('columns', [])
            if len(columns) != 2:
                raise VisualError('Comparison layout requires exactly two columns')
            for c_index, section in enumerate(columns):
                x = .65 + c_index * 6.35
                text_box(slide, nonempty(section.get('title'), 'column title'), (x, 1.6, 5.7, .55), font, size=23, bold=True)
                text_box(slide, '\n\n'.join(section.get('body', [])), (x, 2.4, 5.7, 3.9), font, size=22)
        elif layout == 'title':
            text_box(slide, '\n\n'.join(body), (.85, 2.2, 11.6, 3.6), font, size=27)
        elif item.get('image') or item.get('chart'):
            if item.get('image') and item.get('chart'):
                raise VisualError('Use separate slides or custom layout for both an image and chart')
            asset_box = (.65, 1.6, 12.0, 4.7) if layout == 'figure' else (.65, 1.6, 8.2, 4.9)
            if item.get('image'):
                image = item['image']
                picture(slide, source_path(sources, image.get('source')), asset_box, image.get('alt_text'))
                assets.append(image['source'])
            else:
                native_chart(slide, item['chart'], sources, asset_box, font)
                assets.append(item['chart']['source'])
            if body:
                if layout == 'figure':
                    raise VisualError('Put full-figure commentary in speaker notes or use the evidence layout')
                text_box(slide, '\n\n'.join(body), (9.2, 1.7, 3.45, 4.7), font, size=21)
        else:
            text_box(slide, '\n\n'.join(body), (.85, 1.7, 11.6, 4.7), font, size=25)
        source_ids = list(dict.fromkeys(assets + [e['source'] for c in claim_ids for e in claims[c].get('evidence', [])]))
        status = 'ILLUSTRATIVE' if spec['evidence_status'] == 'illustrative' else ', '.join(sorted({claims[c]['status'] for c in claim_ids}))
        footer = f'{index + 1} | ' + ('Backup | ' if item.get('backup') else '') + status
        if source_ids: footer += ' | ' + '; '.join(sources[s].get('label', s) for s in source_ids)
        text_box(slide, footer, (.65, 6.95, 12, .3), font, size=10, color='445566')
        evidence_text = '\n'.join(f"[{c}] {claims[c]['status']}: {claims[c]['text']}" for c in claim_ids)
        source_text = '\n'.join(f"{s}: {sources[s].get('label', s)}; SHA256 {sources[s]['sha256']}" for s in source_ids)
        slide.notes_slide.notes_text_frame.text = f'Timing: {duration} seconds\n{notes}\n\nClaims\n{evidence_text}\n\nSources\n{source_text}'
        ledger.append({'slide': index + 1, 'title': title, 'claims': claim_ids, 'sources': source_ids,
                       'seconds': duration, 'backup': item.get('backup', False), 'layout': layout})
    if spec.get('duration_seconds') is not None and total_seconds > number(spec['duration_seconds'], 'talk duration', 1, 86400):
        raise VisualError('Main slides exceed the declared talk duration')
    with output_bundle(destination) as out:
        deck.save(out / 'slides.pptx')
        write_json(out / 'slides-ledger.json', ledger)
        (out / 'speaker-notes.md').write_text('\n\n'.join(f"## {i + 1}. {s['title']}\n\n{s['notes']}" for i, s in enumerate(slides)) + '\n', encoding='utf-8')
        return finalize(out, spec, sources, {'slides': len(slides), 'main_talk_seconds': total_seconds,
                        'font': font, 'editable_text': True, 'native_charts': sum(bool(s.get('chart')) for s in slides),
                        'office_render': 'not_run', 'visual_review': 'pending', 'text_fit': 'heuristic screen only'}, builder=Path(__file__))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('spec'); parser.add_argument('--output', required=True)
    args = parser.parse_args()
    try:
        build(args.spec, args.output)
    except (VisualError, OSError, KeyError, TypeError) as exc:
        parser.exit(2, str(exc) + '\n')
